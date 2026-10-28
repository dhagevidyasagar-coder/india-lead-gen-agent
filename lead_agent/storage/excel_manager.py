import os
from datetime import datetime
from typing import List, Dict, Any
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from lead_agent.templates.pitch_templates import PITCH_TEMPLATES

class ExcelManager:
    """Manages the persistence, deduplication, styling, and daily updates of the Excel database."""

    def __init__(self, file_path: str = "AGMIDI_India_Leads.xlsx"):
        self.file_path = file_path

    def load_existing_identifiers(self) -> set:
        """Returns a set of unique identifiers (phone, clean business name, website) already stored."""
        existing = set()
        if not os.path.exists(self.file_path):
            return existing

        try:
            wb = openpyxl.load_workbook(self.file_path, data_only=True)
            if "All Leads (Pan-India)" in wb.sheetnames:
                ws = wb["All Leads (Pan-India)"]
                # Column 3 is Business Name, 7 is Phone, 10 is Website
                for row in ws.iter_rows(min_row=2, values_only=True):
                    name = str(row[2]).strip().lower() if row and len(row) > 2 and row[2] else ""
                    phone = str(row[6]).strip() if row and len(row) > 6 and row[6] else ""
                    website = str(row[9]).strip().lower() if row and len(row) > 9 and row[9] else ""
                    
                    if name:
                        existing.add(name)
                    if phone and phone != "None":
                        existing.add(phone)
                    if website and website != "None" and "http" in website:
                        existing.add(website)
            wb.close()
        except Exception as e:
            print(f"[Warning] Could not read existing leads for deduplication: {e}")

        return existing

    def get_existing_lead_count(self) -> int:
        if not os.path.exists(self.file_path):
            return 0
        try:
            wb = openpyxl.load_workbook(self.file_path, read_only=True)
            if "All Leads (Pan-India)" in wb.sheetnames:
                ws = wb["All Leads (Pan-India)"]
                count = ws.max_row - 1
                wb.close()
                return max(count, 0)
        except Exception:
            pass
        return 0

    def update_excel(self, new_leads: List[Dict[str, Any]]) -> Dict[str, Any]:
        """Appends new deduplicated leads, formats sheets, and logs daily metrics."""
        existing_keys = self.load_existing_identifiers()
        today_str = datetime.now().strftime("%Y-%m-%d")

        # Deduplicate
        filtered_leads = []
        for lead in new_leads:
            name_key = str(lead.get("business_name", "")).strip().lower()
            phone_key = str(lead.get("phone", "")).strip()
            web_key = str(lead.get("website", "")).strip().lower()

            is_duplicate = False
            if name_key and name_key in existing_keys:
                is_duplicate = True
            elif phone_key and phone_key in existing_keys:
                is_duplicate = True
            elif web_key and web_key in existing_keys and "http" in web_key:
                is_duplicate = True

            if not is_duplicate:
                filtered_leads.append(lead)
                if name_key:
                    existing_keys.add(name_key)
                if phone_key:
                    existing_keys.add(phone_key)
                if web_key:
                    existing_keys.add(web_key)

        # Open or create workbook
        if os.path.exists(self.file_path):
            wb = openpyxl.load_workbook(self.file_path)
        else:
            wb = openpyxl.Workbook()
            # Remove default sheet
            if "Sheet" in wb.sheetnames:
                wb.remove(wb["Sheet"])

        # 1. Main Sheet: All Leads
        ws_leads = self._setup_leads_sheet(wb)
        start_row = ws_leads.max_row + 1
        current_id_num = ws_leads.max_row  # rough offset

        # Append new leads
        for idx, lead in enumerate(filtered_leads, start=1):
            lead_id = f"AGM-{1000 + current_id_num + idx}"
            row_num = start_row + idx - 1
            
            row_data = [
                lead_id,
                lead.get("date_added", today_str),
                lead.get("business_name", ""),
                lead.get("category", "Coaching / EdTech"),
                lead.get("city", "Pan-India"),
                lead.get("state", "India"),
                lead.get("phone", ""),
                lead.get("whatsapp_link", ""),
                lead.get("email", ""),
                lead.get("website", ""),
                lead.get("lead_score", 50),
                lead.get("priority_tier", "Tier B"),
                lead.get("source", "Web Discovery"),
                lead.get("status", "New Lead"),
                lead.get("outreach_stage", "Not Contacted"),
                "Branded App & Student Ads",
                lead.get("description", "")
            ]
            ws_leads.append(row_data)

            # Apply row styling
            self._style_lead_row(ws_leads, row_num, lead.get("lead_score", 50))

        # 2. Daily Log Sheet
        self._record_daily_log(wb, len(filtered_leads))

        # 3. Outreach Pitch Templates Sheet
        self._setup_templates_sheet(wb)

        # Auto-adjust column widths
        for ws in wb.worksheets:
            self._autofit_columns(ws)

        wb.save(self.file_path)
        wb.close()

        total_count = self.get_existing_lead_count()
        return {
            "new_added": len(filtered_leads),
            "skipped_duplicates": len(new_leads) - len(filtered_leads),
            "total_leads_in_file": total_count,
            "file_path": os.path.abspath(self.file_path)
        }

    def _setup_leads_sheet(self, wb: openpyxl.Workbook) -> openpyxl.worksheet.worksheet.Worksheet:
        sheet_name = "All Leads (Pan-India)"
        if sheet_name in wb.sheetnames:
            return wb[sheet_name]

        ws = wb.create_sheet(title=sheet_name, index=0)
        headers = [
            "Lead ID", "Discovery Date", "Business / Creator Name", "Category",
            "City", "State", "Phone (+91)", "WhatsApp Direct Link", "Email Address",
            "Website / Channel", "Lead Score (1-100)", "Priority Tier",
            "Data Source", "Lead Status", "Outreach Stage", "AGMIDI Pitch Offer", "Notes / Description"
        ]
        ws.append(headers)

        # Style header row
        header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_alignment

        ws.row_dimensions[1].height = 28
        ws.freeze_panes = "A2"
        return ws

    def _style_lead_row(self, ws, row_num: int, score: int):
        thin_border = Border(
            left=Side(style="thin", color="E2E8F0"),
            right=Side(style="thin", color="E2E8F0"),
            top=Side(style="thin", color="E2E8F0"),
            bottom=Side(style="thin", color="E2E8F0")
        )

        # Highlight high-score leads with subtle mint green
        row_fill = None
        if score >= 75:
            row_fill = PatternFill(start_color="F0FDF4", end_color="F0FDF4", fill_type="solid")
        elif row_num % 2 == 0:
            row_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

        for col_idx in range(1, 18):
            cell = ws.cell(row=row_num, column=col_idx)
            cell.border = thin_border
            if row_fill:
                cell.fill = row_fill
            cell.font = Font(name="Calibri", size=10)
            
            # Alignments
            if col_idx in [1, 2, 7, 11, 12, 14, 15]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

        ws.row_dimensions[row_num].height = 22

    def _record_daily_log(self, wb: openpyxl.Workbook, added_count: int):
        sheet_name = "Daily Updates & Metrics"
        if sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
        else:
            ws = wb.create_sheet(title=sheet_name, index=1)
            headers = ["Date & Time", "New Leads Discovered", "Agent Status", "Target Focus", "Notes"]
            ws.append(headers)
            
            header_fill = PatternFill(start_color="0F766E", end_color="0F766E", fill_type="solid")
            header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
            for col_idx in range(1, len(headers) + 1):
                cell = ws.cell(row=1, column=col_idx)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center")
            ws.row_dimensions[1].height = 24
            ws.freeze_panes = "A2"

        timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        ws.append([
            timestamp,
            added_count,
            "Success (Auto-Run)",
            "Pan-India Coaching Hubs & Creators",
            f"Added {added_count} verified leads to AGMIDI pipeline"
        ])

    def _setup_templates_sheet(self, wb: openpyxl.Workbook):
        sheet_name = "Outreach Scripts (AGMIDI)"
        if sheet_name in wb.sheetnames:
            return  # Already exists

        ws = wb.create_sheet(title=sheet_name, index=2)
        ws.append(["Template Name", "Target Segment", "Ready-to-Use Message Script"])

        header_fill = PatternFill(start_color="4C1D95", end_color="4C1D95", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        for col in range(1, 4):
            c = ws.cell(row=1, column=col)
            c.fill = header_fill
            c.font = header_font
            c.alignment = Alignment(horizontal="center", vertical="center")

        row = 2
        for key, item in PITCH_TEMPLATES.items():
            content = item.get("script", item.get("body", ""))
            ws.append([item.get("title"), item.get("objective", "Director Pitch"), content])
            ws.cell(row=row, column=3).alignment = Alignment(wrap_text=True, vertical="top")
            ws.row_dimensions[row].height = 120
            row += 1

    def _autofit_columns(self, ws):
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                # Do not base column width on long descriptions or scripts
                if cell.value and cell.column != 17 and cell.column != 3:
                    val_str = str(cell.value)
                    if len(val_str) > max_len and len(val_str) < 50:
                        max_len = len(val_str)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)
